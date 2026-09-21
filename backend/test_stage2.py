import sys
import docx
from app.services.parser import ScriptParser

# Принудительная установка UTF-8 для вывода в консоль Windows
sys.stdout.reconfigure(encoding='utf-8')

def test_parser():
    print("=== ТЕСТ 1: Разбор размеченного текста ===")
    sample_text = """
Секреты Космоса

Сцена 1: Запуск ракеты
[Кадр 1]
Визуал: Огромная ракета стоит на стартовой площадке на закате
Озвучка: Добро пожаловать на канал. Сегодня мы отправляемся в путешествие к далеким звездам.

[Кадр 2]
Визуал: Двигатели ракеты загораются, появляется пламя и клубы дыма
Озвучка: Запуск ракеты — это сочетание непревзойденной инженерной мысли и огромной энергии.

Сцена 2: Открытый космос
[Кадр 1]
Визуал: Вид на Землю с орбиты, синий океан и белые облака
Озвучка: Мы в открытом космосе. Наша планета выглядит как хрупкий голубой шарик.
"""

    res1 = ScriptParser.parse_text_to_script(sample_text)
    print(f"Заголовок: {res1.title}")
    print(f"Всего сцен: {res1.total_scenes}, всего кадров: {res1.total_frames}")
    for sc in res1.scenes:
        print(f"\n  [{sc.title}] (Кадров: {len(sc.frames)})")
        for fr in sc.frames:
            print(f"    - {fr.frame_id}:")
            print(f"        Промпт: {fr.visual_prompt}")
            print(f"        Озвучка: {fr.narration_text}")

    print("\n=== ТЕСТ 2: Разбор DOCX файла ===")
    doc = docx.Document()
    doc.add_heading('История ИИ', 0)
    doc.add_paragraph('Сцена 1: Зарождение')
    doc.add_paragraph('Искусственный интеллект начал свое развитие в середине XX века.')
    doc.add_paragraph('Сцена 2: Наши дни')
    doc.add_paragraph('Сегодня нейросети генерируют видео и помогают ученым совершать открытия.')
    
    doc_path = "test_script.docx"
    doc.save(doc_path)
    
    with open(doc_path, "rb") as f:
        text = ScriptParser.parse_docx_bytes(f.read())
    
    res2 = ScriptParser.parse_text_to_script(text, default_title="История ИИ")
    print(f"Заголовок: {res2.title}")
    print(f"Всего сцен: {res2.total_scenes}, всего кадров: {res2.total_frames}")

    assert res1.total_scenes == 2, "Должно быть 2 сцены в тесте 1"
    assert res1.total_frames == 3, "Должно быть 3 кадра в тесте 1"
    assert res2.total_scenes == 2, "Должно быть 2 сцены в тесте 2"
    print("\n[SUCCESS] Все модульные тесты парсера успешно пройдены!")

if __name__ == "__main__":
    test_parser()
